import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)

def create_guide_docx():
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Palette
    COLOR_PRIMARY = RGBColor(15, 61, 110)    # #0F3D6E Navy
    COLOR_SECONDARY = RGBColor(13, 148, 136) # #0D9488 Teal
    COLOR_ACCENT = RGBColor(217, 119, 6)     # #D97706 Amber
    COLOR_DARK = RGBColor(30, 41, 59)        # #1E293B Slate Dark
    COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B Slate Muted
    COLOR_CODE = RGBColor(37, 99, 235)       # #2563EB Blue

    # 1. Document Title Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_badge = title_p.add_run("TÀI LIỆU KỸ THUẬT & HƯỚNG DẪN VẬN HÀNH CHUYÊN SÂU")
    run_badge.font.name = "Arial"
    run_badge.font.size = Pt(9.5)
    run_badge.font.bold = True
    run_badge.font.color.rgb = COLOR_ACCENT

    h1_p = doc.add_paragraph()
    h1_p.paragraph_format.space_before = Pt(2)
    h1_p.paragraph_format.space_after = Pt(8)
    run_title = h1_p.add_run("HƯỚNG DẪN CÀI ĐẶT & SỬ DỤNG TENCENT BROWSERSKILL\nVÀ HỆ THỐNG DỰ ÁN LOCAL")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Nền tảng Tự Động Hóa Trình Duyệt Thực Tế Cho AI Agent (Cursor, Antigravity, Claude Code, Codex) & Cẩm Nang Tiện Ích Đóng Dấu Bản Quyền, Hệ Thống Pop-up")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_MUTED

    # Meta Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        [("Tác giả / Nền tảng:", "Tencent Open Source & Ninh Bình Digital Team"), ("Phiên bản BrowserSkill CLI:", "v0.3.0 (Protocol v1.3)")],
        [("Môi trường máy chủ local:", "Windows x64 / PowerShell / Python 3.14"), ("Cập nhật mới:", "Watermark đa năng + Modal Pop-up hiện đại")]
    ]
    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            label, val = meta_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            r1 = p.add_run(label + " ")
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = COLOR_PRIMARY
            r2 = p.add_run(val)
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = COLOR_DARK
            set_cell_background(cell, "F1F5F9" if row_idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper functions
    def add_section_heading(text, num=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{num} {text}".strip())
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY

    def add_sub_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY

    def add_body_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_DARK
        return p

    def add_callout(text, title="LƯU Ý QUAN TRỌNG", box_type="info"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        bg_hex = "EFF6FF" if box_type == "info" else ("FEF3C7" if box_type == "warn" else "F0FDF4")
        border_color = COLOR_CODE if box_type == "info" else (COLOR_ACCENT if box_type == "warn" else COLOR_SECONDARY)
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"📌 {title}: ")
        r_title.font.name = "Arial"
        r_title.font.size = Pt(9.5)
        r_title.font.bold = True
        r_title.font.color.rgb = border_color
        r_body = p.add_run(text)
        r_body.font.name = "Arial"
        r_body.font.size = Pt(9.5)
        r_body.font.color.rgb = COLOR_DARK

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(4)

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "0F172A") # Dark slate
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(code_text)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(226, 232, 240)

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(4)

    # =========================================================================
    # PHẦN 1
    # =========================================================================
    add_section_heading("TỔNG QUAN VỀ TENCENT BROWSERSKILL", "PHẦN 1:")
    add_body_p("Tencent BrowserSkill là giải pháp tự động hóa trình duyệt web thế hệ mới dành cho các hệ thống AI Agent (Cursor, Antigravity, Claude Code, Codex, OpenClaw...). Khác với Selenium hay Puppeteer truyền thống thường mở một trình duyệt trống không có đăng nhập, BrowserSkill kết nối trực tiếp vào trình duyệt Chrome/Edge thực tế bạn đang dùng.")

    add_sub_heading("4 Ưu Điểm Đột Phá Của BrowserSkill:")
    add_body_p("AI tận dụng các phiên đăng nhập tài khoản có sẵn (Google, Facebook, Zalo, Portal quản trị...) mà bạn không cần phải chia sẻ mật khẩu hay tạo tài khoản phụ.", "1. Tái sử dụng phiên đăng nhập thật (Real Login State): ")
    add_body_p("Mọi tác vụ của AI được chạy trong một cửa sổ làm việc riêng (Agent Window). Bạn vẫn có thể tiếp tục lướt web, làm việc trên các tab khác mà không hề bị chuột hay bàn phím của AI cướp tiêu điểm.", "2. Không làm gián đoạn người dùng (Uninterrupted Work): ")
    add_body_p("Khi AI gặp mã xác thực Captcha, xác thực OTP 2 lớp hoặc giao dịch tiền bạc, nó sẽ gọi lệnh request-help để bạn vào xử lý rồi tiếp tục tự động chạy.", "3. Cơ chế Human-in-the-Loop thông minh: ")
    add_body_p("Trang bị công nghệ chụp ảnh màn hình cuộn toàn trang (Full-page screenshot) và sơ đồ phần tử ngữ nghĩa VOM/Aria-snapshot giúp AI nhận diện nút bấm cực kỳ chính xác.", "4. Chụp toàn trang & Nhận diện ngữ nghĩa: ")

    # =========================================================================
    # PHẦN 2
    # =========================================================================
    add_section_heading("QUY TRÌNH CÀI ĐẶT & KẾT NỐI (ĐÃ THỰC HIỆN TRÊN MÁY BẠN)", "PHẦN 2:")
    add_body_p("Hệ thống BrowserSkill bao gồm 2 thành phần chính: bsk CLI / Daemon chạy ngầm trên máy tính và Browser Extension cài vào trình duyệt web.")

    add_sub_heading("Bước 1: Cài đặt CLI bsk (Đã hoàn thành 100%)")
    add_body_p("CLI bsk đã được tải và cài đặt thành công vào thư mục người dùng:")
    add_code_block("Đường dẫn thực thi: C:\\Users\\Admin\\.local\\bin\\bsk.exe\nPhiên bản hiện tại: bsk v0.3.0 (Protocol version 1.3)")

    add_sub_heading("Bước 2: Cài đặt Agent Skill (Đã tích hợp vào các AI Agent)")
    add_body_p("Skill browser-skill đã được nạp sẵn vào các môi trường:")
    add_body_p("• C:\\Users\\Admin\\.agents\\skills\\browser-skill\\SKILL.md (Dành cho Codex / Antigravity)")
    add_body_p("• C:\\Users\\Admin\\.cursor\\skills\\browser-skill\\SKILL.md (Dành cho Cursor)")
    add_body_p("• C:\\Users\\Admin\\.claude\\skills\\browser-skill\\SKILL.md (Dành cho Claude Code)")
    add_body_p("• C:\\Users\\Admin\\.gemini\\config\\skills\\browser-skill\\SKILL.md (Dành cho Antigravity Global)")

    add_sub_heading("Bước 3: Cài đặt Extension trên trình duyệt (Bạn thực hiện một lần)")
    add_body_p("Để trình duyệt kết nối được với AI, bạn chỉ cần mở trình duyệt và cài đặt tiện ích từ Store chính thức:")
    add_body_p("• Google Chrome: https://chromewebstore.google.com/detail/hhcmgoofomhgciiibhipgmgkgnoenaoi", "🔗 ")
    add_body_p("• Microsoft Edge: https://microsoftedge.microsoft.com/addons/detail/browserskill/emacgiaaaiojkkpkddmmdfhmokgmnikg", "🔗 ")
    add_body_p("Sau khi bấm 'Thêm vào Chrome/Edge', extension sẽ hiển thị icon BrowserSkill trên thanh công cụ và tự động kết nối tới Daemon đang chạy ở cổng 52800.")

    add_sub_heading("Bước 4: Khởi động Daemon nền")
    add_body_p("Tôi đã tạo sẵn file khởi động 1-click trong thư mục dự án:")
    add_body_p("👉 File: c:\\Users\\Admin\\Downloads\\Allinone\\start_browserskill.bat")
    add_body_p("Chỉ cần nhấp đúp vào file này, Daemon sẽ tự khởi động và lắng nghe kết nối WebSocket tại ws://127.0.0.1:52800.")

    add_callout("Khi bạn muốn AI điều khiển trình duyệt, hãy đảm bảo cửa sổ start_browserskill.bat đang mở hoặc Daemon đang chạy ngầm trên máy tính.", "Mẹo vận hành", "info")

    # =========================================================================
    # PHẦN 3
    # =========================================================================
    add_section_heading("HƯỚNG DẪN SỬ DỤNG CHO AI AGENT & LỆNH CLI", "PHẦN 3:")
    add_body_p("Sau khi Extension đã kết nối, bạn có thể tương tác bằng 2 cách rất thuận tiện:")

    add_sub_heading("Cách 1: Ra lệnh tự nhiên cho AI (Khuyên dùng)")
    add_body_p("Bạn chỉ cần chat trực tiếp trong cửa sổ lập trình với AI, AI sẽ tự động gọi các công cụ của BrowserSkill:")
    add_body_p("• 'Dùng browser-skill mở trang https://example.com và chụp ảnh toàn trang lưu thành ketqua.png'")
    add_body_p("• 'Dùng browser-skill vào website sản phẩm của tôi, kiểm tra tính năng đặt hàng xem có lỗi gì không'")
    add_body_p("• 'Dùng browser-skill tìm kiếm từ khóa X trên Google và tóm tắt 3 kết quả đầu tiên'")

    add_sub_heading("Cách 2: Tự thao tác qua lệnh CLI trong Terminal / PowerShell")
    add_body_p("Dưới đây là bảng tra cứu các lệnh quan trọng nhất của bsk:")

    # Table of CLI commands
    cmd_table = doc.add_table(rows=8, cols=3)
    cmd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cmd_table.autofit = False

    cmd_headers = ["Lệnh CLI", "Ý nghĩa / Tác dụng", "Ví dụ áp dụng"]
    for i, h in enumerate(cmd_headers):
        cell = cmd_table.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "0F3D6E")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    cmd_data = [
        ("bsk status", "Kiểm tra tình trạng daemon và số lượng browser kết nối", "bsk status --json"),
        ("bsk doctor", "Chẩn đoán toàn diện sức khỏe hệ thống và gợi ý khắc phục", "bsk doctor"),
        ("bsk session start", "Bắt đầu một phiên làm việc mới (trả về session ID)", "bsk session start --json"),
        ("bsk navigate <url>", "Mở một địa chỉ trang web trong phiên làm việc", "bsk navigate https://google.com --session <id>"),
        ("bsk observe", "Quét giao diện trang web, gán mã @e1, @e2 cho các nút/ô", "bsk observe --session <id>"),
        ("bsk click <@ref>", "Nhấp chuột vào một phần tử trên trang", "bsk click @e3 --session <id>"),
        ("bsk fill <@ref> --value", "Điền văn bản vào ô nhập liệu", "bsk fill @e2 --value 'Ninh Bình' --session <id>")
    ]

    for row_idx, data in enumerate(cmd_data, start=1):
        for col_idx, text in enumerate(data):
            cell = cmd_table.rows[row_idx].cells[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.name = "Consolas"
                r.font.bold = True
                r.font.color.rgb = COLOR_CODE
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # =========================================================================
    # PHẦN 4
    # =========================================================================
    add_section_heading("HƯỚNG DẪN CÁC TÍNH NĂNG MỚI CỦA DỰ ÁN LOCAL", "PHẦN 4:")
    add_body_p("Bên cạnh BrowserSkill, dự án All in One Ninh Bình Digital đã được nâng cấp toàn diện các tính năng quan trọng:")

    add_sub_heading("1. Khởi động Web Server Local")
    add_body_p("• Địa chỉ truy cập chính: http://localhost:5000/index.html")
    add_body_p("• Trực tiếp Soạn Thảo Văn Bản Hành Chính: http://localhost:5000/index.html#vanban")
    add_body_p("• File khởi động nhanh: chay_localhost.bat hoặc start_server.bat")

    add_sub_heading("2. Công cụ Đóng Dấu Bản Quyền Ảnh (Watermark Pro)")
    add_body_p("Công cụ đóng dấu đã được nâng cấp toàn diện để bảo vệ hình ảnh đặc sản, homestay:")
    add_body_p("• Chọn màu sắc watermark: Hỗ trợ bộ chọn màu Color Picker tự do và dãy nút màu chuẩn (Trắng, Đen, Đỏ, Vàng kim, Xanh dương). Kèm tùy chọn viền bóng chữ giúp chữ luôn đọc rõ dù trên nền ảnh sáng hay tối.")
    add_body_p("• Điều chỉnh độ mờ nhạt (Opacity): Thanh trượt trực quan từ 5% (rất mờ tinh tế) đến 100% (rõ nét), cập nhật trực tiếp trên khung xem trước.")
    add_body_p("• Tải lên Watermark Logo: Hỗ trợ nạp file logo PNG trong suốt, điều chỉnh kích cỡ logo theo % so với ảnh.")
    add_body_p("• 6 vị trí đóng dấu: 4 góc, chính giữa hoặc lặp chéo toàn màn hình chống trộm ảnh.")
    add_body_p("• Xuất file đa định dạng: Tải về định dạng JPG chất lượng cao hoặc PNG nguyên bản không giảm độ nét.")

    add_sub_heading("3. Hệ Thống Pop-up Modal Thay Thế Toàn Bộ Alert / Confirm")
    add_body_p("Toàn bộ các hộp thoại alert() và confirm() mặc định của trình duyệt (vốn thô cứng và có dòng chữ 'localhost:5000 cho biết') đã được thay thế 100% bằng Pop-up Modal thiết kế riêng:")
    add_body_p("• Giao diện mờ nền kính (Backdrop blur), bo góc 18px mềm mại, icon trạng thái sinh động.")
    add_body_p("• Hỗ trợ phím tắt tiện dụng: Bấm Enter để xác nhận, bấm Esc để đóng/hủy hộp thoại.")

    add_sub_heading("4. Chuyển Nhóm 5 (Chấm Điểm AI) Sang Trạng Thái 'Đang Phát Triển'")
    add_body_p("Thẻ Chấm Điểm Tự Động & Cổng Nộp Bài Học Sinh tại Nhóm 5 đã được gắn huy hiệu '🚧 Đang phát triển'. Khi nhấp vào, hệ thống hiển thị pop-up thông báo lịch sự về lộ trình ra mắt thay vì điều hướng vào giao diện chưa hoàn thiện.")

    # =========================================================================
    # PHẦN 5
    # =========================================================================
    add_section_heading("XỬ LÝ SỰ CỐ THƯỜNG GẶP (TROUBLESHOOTING)", "PHẦN 5:")
    
    add_body_p("1. Kiểm tra bạn đã bấm mở file start_browserskill.bat chưa. 2. Kiểm tra Extension trên Chrome/Edge đã được bật hay chưa.", "Lỗi '0 browsers connected': ")
    add_body_p("Mở Command Prompt chạy lệnh: netstat -ano | findstr :52800 để kiểm tra PID, sau đó tắt tiến trình cũ và mở lại start_browserskill.bat.", "Lỗi port 52800 bị chiếm dụng: ")
    add_body_p("Luôn đảm bảo câu lệnh kết thúc bằng 'bsk session stop <id>' để trả lại quyền điều khiển bình thường cho trình duyệt.", "Trình duyệt bị treo Agent Window: ")

    # Footer note
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot = p_foot.add_run("Tài liệu được xuất tự động bởi Ninh Bình Digital Assistant • 2026")
    r_foot.font.size = Pt(8.5)
    r_foot.font.italic = True
    r_foot.font.color.rgb = COLOR_MUTED

    # Save to Downloads and local project
    download_dir = os.path.join(os.environ.get("USERPROFILE", "C:\\Users\\Admin"), "Downloads")
    target_download_path = os.path.join(download_dir, "HUONG_DAN_SU_DUNG_BROWSERSKILL_VA_DU_AN.docx")
    target_local_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "HUONG_DAN_SU_DUNG_BROWSERSKILL_VA_DU_AN.docx")

    doc.save(target_download_path)
    doc.save(target_local_path)
    print(f"SUCCESS: Saved to {target_download_path}")
    print(f"SUCCESS: Saved to {target_local_path}")

if __name__ == "__main__":
    create_guide_docx()
